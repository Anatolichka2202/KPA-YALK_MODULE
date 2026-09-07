from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\qt_repos_2\liborbita")
OUT = ROOT / "output" / "documents" / "Вопросник_по_проверке_УБСИ_по_ТУ.docx"


NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F6F7F9"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
GREEN = "217346"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row) -> None:
    repeat_table_header(row)


def set_font(run, name="Arial", size=10.5, bold=None, color="000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_run_text(paragraph, text, *, bold=False, italic=False, size=10.5, color="000000"):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    run.italic = italic
    return run


def style_paragraph(paragraph, before=0, after=5, line=1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_answer_lines(doc: Document, lines=2) -> None:
    for _ in range(lines):
        p = doc.add_paragraph()
        style_paragraph(p, after=2, line=1.0)
        add_run_text(p, "Ответ  __________________________________________________________________________________", size=9.5, color=MID_GRAY)


def add_question(doc: Document, number: str, reference: str, text: str, lines=2) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, before=6, after=2, line=1.12)
    p.paragraph_format.keep_with_next = True
    add_run_text(p, f"{number}  ", bold=True, size=10.5, color=NAVY)
    add_run_text(p, f"{reference}  ", bold=True, size=9.5, color=MID_GRAY)
    add_run_text(p, text, size=10.5)
    add_answer_lines(doc, lines)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=3, line=1.08)
        add_run_text(p, item, size=10.5)


def set_col_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def format_table(table, header_fill=NAVY, font_size=9.2, center_cols: set[int] | None = None) -> None:
    center_cols = center_cols or set()
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_table_header(table.rows[0])
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                set_cell_shading(cell, header_fill)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in center_cols or ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                style_paragraph(p, after=0, line=1.02)
                for run in p.runs:
                    set_font(run, size=font_size, bold=(ri == 0), color=("FFFFFF" if ri == 0 else "000000"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_font(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for sty_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    sty = styles[sty_name]
    sty.font.name = "Arial"
    sty._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    sty._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    sty.font.size = Pt(size)
    sty.font.bold = True
    sty.font.color.rgb = RGBColor(0, 0, 0)
    sty.paragraph_format.space_before = Pt(10)
    sty.paragraph_format.space_after = Pt(5)
    sty.paragraph_format.keep_with_next = True

for sec in doc.sections:
    add_page_number(sec.footer.paragraphs[0])

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
style_paragraph(p, before=8, after=9, line=1.0)
p.add_run("Вопросник по проверке УБСИ по ТУ")

p = doc.add_paragraph()
style_paragraph(p, after=13)
add_run_text(p, "Согласование режимов приемки и производства в стендовом ПО", size=13, color=MID_GRAY)

meta = doc.add_table(rows=4, cols=2)
meta_data = [
    ("Изделие", "УБСИ ЛВРМ.468157.002"),
    ("Нормативный источник", "УБСИ ТУ проект, редакция файла от 13.04.2026"),
    ("Получатели", "Производство и главный инженер"),
    ("Назначение", "Получить решения для полного сценария, интерфейса и протоколов"),
]
for row, (a, b) in zip(meta.rows, meta_data):
    row.cells[0].text = a
    row.cells[1].text = b
format_table(meta, header_fill=PALE_BLUE, font_size=10, center_cols=set())
set_col_widths(meta, [4.2, 12.2])
for row in meta.rows:
    set_cell_shading(row.cells[0], PALE_BLUE)
    for run in row.cells[0].paragraphs[0].runs:
        set_font(run, size=10, bold=True, color="000000")
    for run in row.cells[1].paragraphs[0].runs:
        set_font(run, size=10, color="000000")

doc.add_paragraph()
p = doc.add_paragraph()
style_paragraph(p, after=7, line=1.14)
add_run_text(p, "Цель документа. ", bold=True)
add_run_text(
    p,
    "Зафиксировать фактический производственный процесс и получить нормативные решения там, где проект ТУ задает требование, но не раскрывает алгоритм подпрограммы проверки. Ответы используются для двух режимов, работающих на одном сценарном движке.",
)

add_bullets(
    doc,
    [
        "Приемо-сдаточный режим: минимум информации на экране оператора и полный набор проверок, отнесенных проектом ТУ к п. 5.6.",
        "Производственный режим: многоэтапная проверка с повторными запусками, регистрацией замен и поддержкой методов 5.7 и 5.8.",
        "ЯВП проверяется. Требуется утвердить методику, частотные точки, внешнюю кроссировку и состав средств измерений по пп. 1.1.4.7, 1.1.4.8 и 1.1.4.14.",
        "Эталон 6,2 ± 0,03 В проверяется по значению в выходном потоке телеметрической информации согласно п. 1.1.4.9. Требуется определить параметр протокола Орбита и схему чтения через БСИ.",
        "Форма отчета по ТУ пока не утверждена. Вопросник запрашивает состав и шаблон приемо-сдаточного протокола и производственной ведомости.",
    ],
)

p = doc.add_paragraph()
style_paragraph(p, before=8, after=3)
add_run_text(p, "Как отвечать", bold=True, size=11)
add_bullets(
    doc,
    [
        "Производству: описать фактический порядок работы, включая ручные действия и отличия между этапами.",
        "Главному инженеру: дать решение по каждому нормативному вопросу, указать применяемую редакцию документа и, при наличии, приложить методику или утвержденный шаблон.",
        "Если ответ зависит от комплектации БСИ, кабеля или оснастки, указать обозначение и ревизию.",
    ],
)

doc.add_page_break()

# Reference map
doc.add_heading("Карта требований проекта ТУ", level=1)
p = doc.add_paragraph()
style_paragraph(p, after=7)
add_run_text(
    p,
    "Пункт 5.6 прямо относит к проверке функционирования пп. 1.1.4.1–1.1.4.3, 1.1.4.5–1.1.4.11, 1.1.4.13 и 1.1.4.14. Пункты 1.1.4.4 и 1.1.4.12 проверяются отдельно методом 5.5.",
)

requirements = [
    ("1.1.4.1", "Потенциометрические сигналы 0…6,2 ± 0,2 В, не менее 60 каналов; контактные сигналы, не менее 30 каналов; термосопротивления 0…240 Ом по четырехпроводной схеме, не менее 30 каналов; пьезоэлектрические сигналы 2…2000 Гц, не менее 16 каналов."),
    ("1.1.4.2", "Питание датчиков 6,2 ± 0,2 В при токе до 350 мА; токовая защита от 450 мА и автоматическое восстановление."),
    ("1.1.4.3", "Работа при 24…35 В; сохранение работоспособности при 19 В не менее 5 мин и 37 В не менее 1 мин."),
    ("1.1.4.5", "Ток потребления не более 400 мА при напряжении питания 24…35 В."),
    ("1.1.4.6", "Соединительная линия термодатчик–УБСИ длиной до 50 м."),
    ("1.1.4.7", "Неравномерность АЧХ пьезоэлектрических цепей по таблице 1.1; затухание на 2Fв и выше не менее 20–30 дБ."),
    ("1.1.4.8", "Коэффициенты усиления 0,25; 0,5; 1; 2; 4; 8; 32 мВ/пКл, выбор внешней кроссировкой."),
    ("1.1.4.9", "Эталонное напряжение 6,2 ± 0,03 В с выдачей значения в выходной поток телеметрической информации."),
    ("1.1.4.10", "Идентификация обрыва любого функционального аналогового входа напряжением ниже 0 В."),
    ("1.1.4.11", "Защита входов при перегрузке ±12 В с сохранением работоспособности остальных входов."),
    ("1.1.4.13", "Готовность к работе после подачи питания не более 30 с."),
    ("1.1.4.14", "Погрешность: медленные аналоговые ±0,5 % шкалы, при крайних температуре и питании не более 0,75 %; температурные ±0,5 %; быстроменяющиеся ±7 %; входной ток канала не более 2 мкА."),
    ("5.7", "Работа при минус 50 °C и плюс 50 °C; проверки выполняются подпрограммой УБСИ с учетом п. 1.1.4.14."),
    ("5.8", "Прочность при предельных температурах: минус 60 °C 24 ч и плюс 90 °C 6 ч; после возврата в нормальные условия повторяют проверку по 5.6.2–5.6.9."),
]
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Пункт ТУ"
table.rows[0].cells[1].text = "Требование для согласования"
for ref, req in requirements:
    cells = table.add_row().cells
    cells[0].text = ref
    cells[1].text = req
format_table(table, font_size=8.7, center_cols={0})
set_col_widths(table, [2.6, 13.8])

doc.add_page_break()

# Production questions
doc.add_heading("Вопросы производству", level=1)
p = doc.add_paragraph()
style_paragraph(p, after=8)
add_run_text(p, "Просьба отвечать по фактическому технологическому процессу. Если практика отличается от проекта ТУ, укажите оба варианта и причину отличия.")

doc.add_heading("Этапы и состав проверки", level=2)
prod_questions = [
    ("П 1", "пп. 1.1.3.1, 1.1.3.2; 5.6–5.8", "Сколько фактических этапов проходит УБСИ от изготовления ячеек до приемки готового блока? Укажите названия, порядок и состояние изделия на каждом этапе."),
    ("П 2", "п. 5.6", "На каком этапе выполняется полный приемо-сдаточный прогон по п. 5.6: до заливки, после заливки, после климатических воздействий, перед окончательной сдачей или несколько раз?"),
    ("П 3", "пп. 5.7, 5.8", "Какие проверки реально выполняются до и после рабочих и предельных температур? Совпадает ли повторный прогон с полным п. 5.6 или используется сокращенный набор?"),
    ("П 4", "пп. 1.1.6.1–1.1.6.6", "После каких механических воздействий производство повторяет электрическую проверку УБСИ? Какие параметры обязательны после вибрации и ударов?"),
    ("П 5", "пп. 1.1.4.1–1.1.4.14", "Какие проверки нужны для отдельной ячейки до установки в УБСИ, а какие имеют смысл только в составе собранного УБСИ и БСИ?"),
    ("П 6", "п. 5.6.4", "Нужен ли между длинными этапами короткий контрольный прогон? Если да, перечислите минимальные параметры, достаточные для допуска к следующей операции."),
]
for q in prod_questions:
    add_question(doc, *q, lines=2)

doc.add_heading("Параметры и ручные операции", level=2)
prod_questions_2 = [
    ("П 7", "пп. 1.1.4.1, 1.1.4.14", "ЯЛК: какие точки должны проходить все рабочие каналы в производстве? Подтвердите точки 0; 3,1; 6,2 В, перечень 80 адресов и необходимость отдельных пороговых точек 1,0 и 2,4 В."),
    ("П 8", "пп. 1.1.4.1, 1.1.4.10", "Контактные состояния и обрыв: как производство сейчас формирует состояния 0 и 1 и обрыв, какие признаки наблюдает и по скольким каналам?"),
    ("П 9", "п. 1.1.4.11", "Перегрузка ±12 В: выполняется ли она на каждом УБСИ или выборочно? Какова длительность воздействия, какие каналы перегружают и какие остальные каналы контролируют одновременно?"),
    ("П 10", "пп. 1.1.4.1, 1.1.4.6, 1.1.4.14", "ЯТП: подтвердите общий ручной магазин Р4831 на X123, точки 0; 120; 240 Ом, проверку всех 30 каналов и необходимость эквивалента линии термодатчика длиной 50 м."),
    ("П 11", "пп. 1.1.4.7, 1.1.4.8, 1.1.4.14", "ЯВП: какие операции уже выполняет производство? Укажите каналы, коэффициенты усиления, частоты, амплитуды, внешние кроссировки, приборы и критерий результата."),
    ("П 12", "п. 1.1.4.2", "Питание датчиков: где подключают нагрузку 350 мА и перегрузку 450 мА, сколько выходов проверяют и какое действие оператора нужно автоматизировать или подтвердить вручную?"),
    ("П 13", "пп. 1.1.4.3, 1.1.4.5, 1.1.4.13", "Питание УБСИ: на каких напряжениях измеряют ток и готовность? Подтвердите, должен ли ток ≤400 мА измеряться отдельно при 24, 27 и 35 В или достаточно иной последовательности."),
    ("П 14", "п. 1.1.4.14", "Входной ток ≤2 мкА: выполняется ли измерение на производстве? Укажите разъем, контакты, рассечку, прибор, число проверяемых каналов и допустимость ручного ввода результата."),
    ("П 15", "п. 1.1.4.9", "Эталон 6,2 ±0,03 В: видит ли производство этот параметр в штатной телеметрии Орбита через БСИ? Укажите название параметра, адрес и используемую программу."),
]
for q in prod_questions_2:
    add_question(doc, *q, lines=2)

doc.add_heading("Повторы и регистрация результата", level=2)
prod_questions_3 = [
    ("П 16", "пп. 5.6.4–5.6.7", "Что производство делает при первом результате НЕ НОРМА: проверяет схему и повторяет весь прогон, только отказавший этап или отдельные каналы?"),
    ("П 17", "п. 5.6.5", "Как отличают неисправность изделия от ошибки стенда, связи, оснастки или действия оператора? Какие ошибки не должны считаться браком УБСИ?"),
    ("П 18", "пп. 1.1.3.2, 5.6.6", "Какие данные обязательно записываются в технологический паспорт: изделие, ячейки и их серийные номера, этап, приборы, оператор, дата, результаты, причины повторов и замены?"),
    ("П 19", "пп. 5.6–5.8", "При замене ЯЛК, ЯТП или ЯВП какие ранее пройденные этапы аннулируются и какие проверки новая ячейка должна пройти заново?"),
    ("П 20", "пп. 5.6.6, 5.6.7", "Какие два выходных документа нужны производству: краткий приемо-сдаточный протокол и подробная поканальная ведомость? Перечислите обязательные поля каждого документа."),
    ("П 21", "приложение А, таблица А.1", "Какие средства измерений и оснастка реально закреплены за стендом? Укажите тип, заводской номер, срок поверки или аттестации и допустимые аналоги."),
    ("П 22", "п. 5.1", "Кто и как фиксирует температуру, влажность и давление при обычной проверке? Нужен автоматический датчик условий или достаточно ручного ввода?"),
]
for q in prod_questions_3:
    add_question(doc, *q, lines=2)

# Chief engineer questions
doc.add_heading("Вопросы главному инженеру", level=1)
p = doc.add_paragraph()
style_paragraph(p, after=8)
add_run_text(p, "Ниже перечислены решения, без которых ПО не может достоверно объявить полный результат по п. 5.6. Ответ желательно оформить ссылкой на утвержденную методику, схему, таблицу адресов или ревизию документа.")

doc.add_heading("Границы режима приемки", level=2)
chief_1 = [
    ("ГИ 1", "пп. 1.1.1; 5.6", "Подтвердить: программный приемо-сдаточный режим УБСИ выполняет полный объем п. 5.6. Проверки по 5.2–5.5 оформляются отдельно и не входят в этот автоматический прогон, если договором не установлено иное."),
    ("ГИ 2", "пп. 5.6.4, 5.6.7", "Предоставить утвержденную программу или алгоритм подпрограммы проверки УБСИ, на которую ссылается проект ТУ. Сам п. 5.6 не задает последовательность воздействий, число отсчетов и правила расчета."),
    ("ГИ 3", "пп. 5.6.5, 5.6.7", "Утвердить правила результата: НОРМА, НЕ НОРМА, ОШИБКА СТЕНДА и НЕПОЛНАЯ ПРОВЕРКА; определить допустимое число повторов и область повторного запуска."),
    ("ГИ 4", "п. 5.6.6", "Утвердить форму приемо-сдаточного протокола. Какие реквизиты, подписи, таблицы и сроки хранения обязательны? Допускается ли электронный HTML или PDF с неизменяемым идентификатором запуска?"),
]
for q in chief_1:
    add_question(doc, *q, lines=2)

doc.add_heading("Методики электрических проверок", level=2)
chief_2 = [
    ("ГИ 5", "пп. 1.1.4.1, 1.1.4.14", "Утвердить карту каналов ЯЛК: перечень рабочих адресов, соответствие физическим входам и число каналов, достаточное для подтверждения требований о не менее чем 60 потенциометрических и 30 контактных каналах."),
    ("ГИ 6", "п. 1.1.4.1", "Для контактных входов утвердить способ проверки границ: напряжением 1,0 и 2,4 В, сопротивлением 5 и 100 кОм или обоими способами. Указать, какие значения считаются переходной неопределенной зоной."),
    ("ГИ 7", "пп. 1.1.4.1, 1.1.4.6, 1.1.4.14", "Утвердить методику ЯТП: точки сопротивления, допуск ±0,5 % шкалы, число отсчетов и способ подтверждения работы с линией термодатчика длиной до 50 м."),
    ("ГИ 8", "пп. 1.1.4.7, 1.1.4.8, 1.1.4.14; таблица 1.1", "Утвердить методику ЯВП: Fн и Fв, набор частот в каждой области таблицы 1.1, амплитуду или заряд, коэффициенты усиления, внешнюю кроссировку, число каналов и расчет погрешности ±7 %."),
    ("ГИ 9", "п. 1.1.4.7", "Уточнить критерий затухания на 2Fв и выше: проект ТУ указывает диапазон 20–30 дБ. Какое конкретное минимальное значение использовать для вердикта?"),
    ("ГИ 10", "п. 1.1.4.2", "Утвердить схему проверки выходов 6,2 ±0,2 В при нагрузке до 350 мА и защиты от 450 мА: точки подключения, число выходов, длительность перегрузки, допустимое падение напряжения и время восстановления."),
    ("ГИ 11", "пп. 1.1.4.3, 1.1.4.5", "Утвердить последовательность 24, 27 и 35 В и измерение тока ≤400 мА. Требуется ли отдельный токовый результат при каждом напряжении диапазона?"),
    ("ГИ 12", "п. 1.1.4.3", "Утвердить критерий сохранения работоспособности при 19 В в течение 5 мин и 37 В в течение 1 мин: проверять телеметрию во время воздействия, выполнять полный функциональный цикл после выдержки или применять оба контроля?"),
    ("ГИ 13", "п. 1.1.4.10", "Уточнить охват проверки обрыва: какие входы относятся к любому функциональному аналоговому входу, какое отрицательное напряжение и длительность применять, какой код или признак должен появиться в телеметрии."),
    ("ГИ 14", "п. 1.1.4.11", "Утвердить проверку ±12 В: длительность, последовательность полярностей, перечень перегружаемых входов и критерий сохранения работоспособности остальных каналов."),
    ("ГИ 15", "п. 1.1.4.14", "Утвердить метод измерения входного тока ≤2 мкА: точки рассечки, прибор, рабочее напряжение входа, число каналов и способ учета погрешности средства измерений."),
    ("ГИ 16", "п. 1.1.4.13", "Утвердить событие начала и конца отсчета готовности ≤30 с и минимальный набор признаков, подтверждающий готовность УБСИ."),
]
for q in chief_2:
    add_question(doc, *q, lines=2)

doc.add_heading("Эталон 6,2 В и протокол Орбита", level=2)
chief_3 = [
    ("ГИ 17", "п. 1.1.4.9", "Указать точный параметр протокола Орбита, в котором УБСИ выдает значение эталона 6,2 ±0,03 В: адрес, тип кадра, номер слова или поля, масштаб, единицу измерения, признаки достоверности и период выдачи."),
    ("ГИ 18", "пп. 1.1.4.9; 5.6.1; рисунок А.1", "Подтвердить схему чтения: УБСИ → БСИ → E20 → декодер Орбита либо иной тракт. Является ли БСИ обязательной частью приемо-сдаточной проверки УБСИ по п. 5.6?"),
    ("ГИ 19", "п. 1.1.4.9", "Утвердить критерий: достаточно одного достоверного значения в диапазоне 6,17…6,23 В либо требуется серия значений, медиана, среднее и ограничение разброса."),
    ("ГИ 20", "п. 1.1.4.9", "Подтвердить, что калибровочные слова адаптерного потока ЯЛК не заменяют проверку эталона в выходной телеметрии Орбита."),
]
for q in chief_3:
    add_question(doc, *q, lines=2)

doc.add_heading("Температурные методы и производственные этапы", level=2)
chief_4 = [
    ("ГИ 21", "пп. 5.7.1–5.7.10", "Подтвердить, что при минус 50 °C проверка выполняется при 24 ±0,1 В, при плюс 50 ±2 °C — при 35 ±0,5 В, и что в обоих случаях используется полный функциональный алгоритм п. 5.6 с допуском 0,75 % по п. 1.1.4.14."),
    ("ГИ 22", "пп. 5.8.1–5.8.8", "Подтвердить порядок проверок после минус 60 °C и плюс 90 °C: полный п. 5.6.2–5.6.9 после выдержки в нормальных условиях. Указать допустимое время между стабилизацией и запуском."),
    ("ГИ 23", "пп. 5.7, 5.8", "Подтвердить, что методы 5.7 и 5.8 включаются в производственный многоэтапный режим, а не в обычный приемо-сдаточный экран."),
    ("ГИ 24", "пп. 1.1.3.1, 1.1.3.2", "Утвердить перечень технологических этапов и правило повторной проверки после замены ячейки. Какие результаты прежней ячейки сохраняются и какие этапы новая ячейка проходит заново?"),
]
for q in chief_4:
    add_question(doc, *q, lines=2)

# Landscape matrices
land = doc.add_section(start_type=1)
land.orientation = WD_ORIENT.LANDSCAPE
land.page_width, land.page_height = section.page_height, section.page_width
land.top_margin = Inches(0.55)
land.bottom_margin = Inches(0.55)
land.left_margin = Inches(0.55)
land.right_margin = Inches(0.55)

doc.add_heading("Матрица производственных этапов для заполнения", level=1)
p = doc.add_paragraph()
style_paragraph(p, after=7)
add_run_text(p, "Отметьте обязательные проверки знаком X. Если этап отсутствует, зачеркните столбец и подпишите фактическое название этапа.")

headers = ["Пункт ТУ", "Проверка", "До сборки", "После сборки", "После вибрации", "5.7 рабочие температуры", "5.8 предельные температуры", "Финальная приемка"]
rows = [
    ("1.1.4.1, 1.1.4.14", "ЯЛК аналоговые каналы"),
    ("1.1.4.1", "Контактные пороги"),
    ("1.1.4.10", "Обрыв входов"),
    ("1.1.4.11", "Перегрузка ±12 В"),
    ("1.1.4.1, 1.1.4.6, 1.1.4.14", "ЯТП 0…240 Ом"),
    ("1.1.4.7, 1.1.4.8, 1.1.4.14", "ЯВП АЧХ и коэффициенты"),
    ("1.1.4.2", "Питание датчиков 350 и 450 мА"),
    ("1.1.4.3", "Питание 24…35 В и выдержки 19/37 В"),
    ("1.1.4.5", "Ток потребления ≤400 мА"),
    ("1.1.4.9", "Эталон 6,2 В в Орбите"),
    ("1.1.4.13", "Готовность ≤30 с"),
    ("1.1.4.14", "Входной ток ≤2 мкА"),
    ("5.1", "Климатические условия и приборы"),
    ("5.6.6", "Протокол и ведомость каналов"),
]
matrix = doc.add_table(rows=1, cols=len(headers))
for i, h in enumerate(headers):
    matrix.rows[0].cells[i].text = h
for ref, test in rows:
    cells = matrix.add_row().cells
    cells[0].text = ref
    cells[1].text = test
    for i in range(2, len(headers)):
        cells[i].text = "[   ]"
format_table(matrix, font_size=8.3, center_cols={0, 2, 3, 4, 5, 6, 7})
set_col_widths(matrix, [3.2, 5.3, 2.3, 2.5, 2.8, 3.5, 3.6, 3.0])

doc.add_paragraph()
doc.add_heading("Матрица автоматизации", level=2)
headers2 = ["Проверка", "Автоматически", "Ручное подтверждение", "Нужна новая оснастка", "Не выполнять на этом этапе", "Комментарий"]
automation_rows = [
    "ЯЛК и контактные состояния",
    "ЯТП и магазин Р4831",
    "ЯВП",
    "Питание датчиков 350/450 мА",
    "Питание УБСИ и ток",
    "Эталон 6,2 В в Орбите",
    "Входной ток ≤2 мкА",
    "Климатические выдержки 5.7/5.8",
    "Формирование документов",
]
matrix2 = doc.add_table(rows=1, cols=len(headers2))
for i, h in enumerate(headers2):
    matrix2.rows[0].cells[i].text = h
for test in automation_rows:
    cells = matrix2.add_row().cells
    cells[0].text = test
    for i in range(1, 5):
        cells[i].text = "[   ]"
    cells[5].text = ""
format_table(matrix2, font_size=8.7, center_cols={1, 2, 3, 4})
set_col_widths(matrix2, [5.1, 3.0, 3.8, 3.8, 4.1, 7.0])

# Final sign-off page in portrait
last = doc.add_section(start_type=1)
last.orientation = WD_ORIENT.PORTRAIT
last.page_width, last.page_height = section.page_width, section.page_height
last.top_margin = Inches(0.65)
last.bottom_margin = Inches(0.65)
last.left_margin = Inches(0.72)
last.right_margin = Inches(0.72)

doc.add_heading("Лист согласования", level=1)
p = doc.add_paragraph()
style_paragraph(p, after=12)
add_run_text(p, "После заполнения вопросы с разными ответами переносятся в перечень открытых решений. Программная реализация полного п. 5.6 начинается после утверждения методик ЯВП, питания датчиков, входного тока и параметра эталона в протоколе Орбита.")

sign = doc.add_table(rows=5, cols=4)
headers3 = ["Роль", "Фамилия и инициалы", "Подпись", "Дата"]
for i, h in enumerate(headers3):
    sign.rows[0].cells[i].text = h
for ri, role in enumerate(["Производство", "ОТК", "Главный инженер", "Разработчик ПО"], start=1):
    sign.rows[ri].cells[0].text = role
format_table(sign, font_size=9.5, center_cols={0, 2, 3})
set_col_widths(sign, [4.0, 6.0, 3.2, 3.2])
for row in sign.rows[1:]:
    for cell in row.cells:
        cell.add_paragraph("\n")

doc.add_paragraph()
p = doc.add_paragraph()
style_paragraph(p, before=10, after=4)
add_run_text(p, "Нормативная ссылка", bold=True)
p = doc.add_paragraph()
style_paragraph(p, after=3)
add_run_text(p, "УБСИ ТУ проект, файл УБСИ_ТУ_Проект.docx, редакция файла от 13.04.2026. Обозначение изделия в тексте проекта ТУ: ЛВРМ.468157.002.", size=9.5, color=MID_GRAY)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
